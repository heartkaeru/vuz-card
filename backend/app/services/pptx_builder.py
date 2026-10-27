import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from app.data.mock_universities import MOCK_UNIVERSITIES_DB

# Директория для сохранения сгенерированных презентаций
STORAGE_DIR = Path("d:/Dev/vuz-card/backend/storage/presentations")
STORAGE_DIR.mkdir(parents=True, exist_ok=True)


def build_university_presentation(university_id: int) -> str:
    """
    Генерирует брендированную презентацию .pptx к встрече с руководством ВУЗа.
    Возвращает путь к созданному файлу.
    """
    uni = MOCK_UNIVERSITIES_DB.get(university_id)
    if not uni:
        raise ValueError(f"ВУЗ {university_id} не найден")

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # Формат 16:9
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Цветовая палитра Сбера
    COLOR_GREEN = RGBColor(33, 160, 56)      # Фирменный зеленый
    COLOR_DARK = RGBColor(27, 31, 35)        # Темный текст
    COLOR_GRAY = RGBColor(100, 110, 120)     # Второстепенный текст

    # ==================== СЛАЙД 1: Титульный ====================
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Зеленый акцент
    shape = slide1.shapes.add_shape(1, Inches(0.8), Inches(1.2), Inches(0.15), Inches(2.5)) # Прямоугольник
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_GREEN
    shape.line.color.rgb = COLOR_GREEN

    # Заголовок
    tx_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.0), Inches(11), Inches(2))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "МАТЕРИАЛЫ К ВСТРЕЧЕ С РУКОВОДСТВОМ"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN

    p2 = tf.add_paragraph()
    p2.text = uni["name"]
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_DARK
    p2.space_before = Pt(14)

    # Подзаголовок
    sub_box = slide1.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(10), Inches(2))
    tf_sub = sub_box.text_frame
    p_sub1 = tf_sub.paragraphs[0]
    p_sub1.text = f"Город: {uni['city']}  |  ИНН: {uni['inn']}  |  Сайт: {uni['website']}"
    p_sub1.font.size = Pt(14)
    p_sub1.font.color.rgb = COLOR_GRAY

    p_sub2 = tf_sub.add_paragraph()
    p_sub2.text = f"Ключевые программы: {', '.join(uni.get('active_projects', []))}"
    p_sub2.font.size = Pt(14)
    p_sub2.font.color.rgb = COLOR_DARK
    p_sub2.space_before = Pt(8)

    # ==================== СЛАЙД 2: Руководство и контакты ====================
    slide2 = prs.slides.add_slide(blank_layout)

    # Заголовок слайда 2
    header_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11), Inches(1))
    tf_h = header_box.text_frame
    p_h = tf_h.paragraphs[0]
    p_h.text = f"Руководство: {uni['short_name']}"
    p_h.font.size = Pt(26)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_DARK

    # Карточки должностных лиц
    officials = uni.get("officials", [])
    top_pos = 1.8
    for idx, off in enumerate(officials[:3]):  # Первые 3 ключевых лица
        card_box = slide2.shapes.add_textbox(Inches(0.8), Inches(top_pos), Inches(11.5), Inches(1.5))
        tf_c = card_box.text_frame
        tf_c.word_wrap = True

        p_name = tf_c.paragraphs[0]
        p_name.text = f"{off['fio']} — {off['position']}"
        p_name.font.size = Pt(18)
        p_name.font.bold = True
        p_name.font.color.rgb = COLOR_GREEN

        p_resp = tf_c.add_paragraph()
        p_resp.text = f"Зона ответственности: {off.get('responsibilities', 'Не указана')}"
        p_resp.font.size = Pt(13)
        p_resp.font.color.rgb = COLOR_DARK
        p_resp.space_before = Pt(4)

        p_contacts = tf_c.add_paragraph()
        p_contacts.text = f"Контакты: {off.get('phone', '-')}  |  {off.get('email', '-')}"
        p_contacts.font.size = Pt(12)
        p_contacts.font.color.rgb = COLOR_GRAY
        p_contacts.space_before = Pt(4)

        top_pos += 1.7

    file_name = f"presentation_{university_id}.pptx"
    output_path = STORAGE_DIR / file_name
    prs.save(str(output_path))
    return str(output_path)
